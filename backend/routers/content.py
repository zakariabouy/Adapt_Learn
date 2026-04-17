import logging
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from shared.database import get_pool
from shared.rag import embed_and_store_content, retrieve_relevant_chunks, reembed_all_content
from shared.guardrails import (
    run_input_guardrails,
    run_output_guardrails,
    validate_json_output,
    log_guardrail_event,
)
from routers.auth import get_current_user
from uuid import uuid4, UUID
import datetime
import json
import os
from langchain_core.messages import HumanMessage
from shared.llm import get_rotating_llm

import io
import pdfplumber
from pptx import Presentation

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/content", tags=["Content"])

def get_llm():
    return get_rotating_llm("gemini-2.5-flash")

@router.post("/upload")
async def upload_content(file: UploadFile = File(...), current_user = Depends(get_current_user)):
    # Validate file type
    allowed_extensions = [".md", ".txt", ".pdf", ".pptx", ".ppt"]
    if not any(file.filename.lower().endswith(ext) for ext in allowed_extensions):
        raise HTTPException(status_code=400, detail="Only .md, .txt, .pdf, or .pptx files are allowed.")

    # Check if user is a teacher
    if current_user["role"] != "teacher":
         raise HTTPException(status_code=403, detail="Only teachers can upload content.")

    content = await file.read()
    filename_lower = file.filename.lower()

    if filename_lower.endswith(".pdf"):
        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text_content = "\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to extract text from PDF: {str(e)}")
    elif filename_lower.endswith(".pptx") or filename_lower.endswith(".ppt"):
        try:
            prs = Presentation(io.BytesIO(content))
            slides_text = []
            for i, slide in enumerate(prs.slides, 1):
                parts = [f"--- Slide {i} ---"]
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = paragraph.text.strip()
                            if text:
                                parts.append(text)
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells)
                            if row_text.strip(" |"):
                                parts.append(row_text)
                slides_text.append("\n".join(parts))
            text_content = "\n\n".join(slides_text)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to extract text from PowerPoint: {str(e)}")
    else:
        text_content = content.decode("utf-8")
    
    if not text_content.strip():
         raise HTTPException(status_code=400, detail="File is empty or contains no readable text.")

    # --- Guardrails: sanitize uploaded content before LLM processing ---
    input_check = await run_input_guardrails(
        text_content,
        user_id=current_user["id"],
        endpoint="/content/upload",
    )
    if not input_check["safe"]:
        await log_guardrail_event(
            event_type="prompt_injection",
            severity="critical",
            action_taken="blocked",
            user_id=current_user["id"],
            endpoint="/content/upload",
            input_snippet=text_content[:500],
            details={"issues": input_check["issues"]},
        )
        raise HTTPException(
            status_code=400,
            detail="Uploaded content was flagged by security guardrails. Please review and resubmit.",
        )
    text_content = input_check["sanitized_text"]

    # --- AI Analysis Phase ---
    prompt = f"""
    Analyze the following educational content and extract metadata.
    Content:
    {text_content[:2000]}
    
    Tasks:
    1. Identify the subject (e.g., Biology, Math, History).
    2. Estimate the target grade level (integer).
    3. Generate 5 IRT-calibrated quiz questions. 
       - Each question needs a difficulty (theta) between -3.0 and 3.0.
       - Each question needs 4 options (A, B, C, D) and a correct_id.
       - Include a helpful hint and a clear explanation.
    
    Return a JSON object with:
    {{
      "subject": "string",
      "grade_level": int,
      "topic": "string",
      "questions": [
        {{
          "text": "string",
          "options": [{{"id": "A", "label": "..."}}, ...],
          "correct_id": "A",
          "difficulty": float,
          "hint": "string",
          "explanation": "string"
        }}
      ]
    }}
    """
    
    try:
        response = await get_llm().ainvoke([HumanMessage(content=prompt)])
        # Clean potential markdown backticks
        raw_json = response.content.strip()
        if "```json" in raw_json:
            raw_json = raw_json.split("```json")[1].split("```")[0].strip()
        elif "```" in raw_json:
            raw_json = raw_json.split("```")[1].split("```")[0].strip()
            
        ai_data = json.loads(raw_json)
    except Exception as e:
        logger.warning("AI Content Parsing Error: %s", e)
        ai_data = {
            "subject": "General",
            "grade_level": 5,
            "topic": file.filename,
            "questions": []
        }

    pool = await get_pool()
    content_id = uuid4()

    # Clamp grade_level to primary-school range [1, 6] — DB has a CHECK constraint.
    try:
        raw_grade = int(ai_data.get("grade_level") or 3)
    except (TypeError, ValueError):
        raw_grade = 3
    grade_level = max(1, min(6, raw_grade))

    # 1. Save Content Item
    await pool.execute(
        "INSERT INTO content_items (id, title, original_text, teacher_id, created_at, subject, grade_level) VALUES ($1, $2, $3, $4, $5, $6, $7)",
        content_id, file.filename, text_content, current_user["id"], datetime.datetime.now(),
        ai_data.get("subject"), grade_level
    )
    
    # 2. Seed Question Bank
    for q in ai_data.get("questions", []):
        try:
            await pool.execute(
                """INSERT INTO question_bank 
                   (id, content_id, subject, topic, difficulty, question_text, options, correct_id, hint, explanation) 
                   VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10)""",
                str(uuid4()), content_id, ai_data.get("subject"), ai_data.get("topic"), 
                q["difficulty"], q["text"], json.dumps(q["options"]), q["correct_id"], 
                q["hint"], q["explanation"]
            )
        except Exception as e:
            logger.warning("Failed to seed AI question: %s", e)
    
    # 3. RAG: chunk + embed content for semantic retrieval
    rag_meta = {"subject": ai_data.get("subject"), "grade_level": ai_data.get("grade_level")}
    chunks_stored = 0
    try:
        chunks_stored = await embed_and_store_content(content_id, text_content, rag_meta)
    except Exception as e:
        logger.warning("RAG embedding failed for content %s (non-blocking): %s", content_id, e)

    return {
        "id": content_id,
        "title": file.filename,
        "subject": ai_data.get("subject"),
        "questions_generated": len(ai_data.get("questions", [])),
        "rag_chunks": chunks_stored,
        "status": "processed"
    }

@router.get("/list")
async def list_content(current_user = Depends(get_current_user)):
    pool = await get_pool()
    if current_user["role"] == "teacher":
        rows = await pool.fetch(
            "SELECT id, title, subject, grade_level FROM content_items WHERE teacher_id = $1 ORDER BY created_at DESC",
            current_user["id"]
        )
    else:
        rows = await pool.fetch(
            """SELECT ci.id, ci.title, ci.subject, ci.grade_level
               FROM content_items ci
               JOIN teacher_student_link tsl ON ci.teacher_id = tsl.teacher_id
               WHERE tsl.student_id = $1
               ORDER BY ci.created_at DESC""",
            current_user["id"]
        )
    return [dict(row) for row in rows]

@router.get("/teacher/dashboard/stats")
async def get_teacher_stats(current_user = Depends(get_current_user)):
    if current_user["role"] != "teacher":
        raise HTTPException(status_code=403, detail="Forbidden")

    pool = await get_pool()

    student_count = await pool.fetchval(
        "SELECT COUNT(*) FROM teacher_student_link WHERE teacher_id = $1",
        current_user["id"]
    )
    content_count = await pool.fetchval(
        "SELECT COUNT(*) FROM content_items WHERE teacher_id = $1",
        current_user["id"]
    )
    # Students with ability < -0.5 are considered at-risk
    risk_alerts = await pool.fetchval(
        """
        SELECT COUNT(*) FROM learner_profiles lp
        JOIN teacher_student_link tsl ON lp.student_id = tsl.student_id
        WHERE tsl.teacher_id = $1
          AND (lp.profile_data->>'ability_estimate')::float < -0.5
        """,
        current_user["id"]
    )
    # Sessions started in last hour with no end time
    active_sessions = await pool.fetchval(
        "SELECT COUNT(*) FROM sessions WHERE started_at > NOW() - INTERVAL '1 hour' AND ended_at IS NULL"
    )

    return {
        "student_count": student_count or 0,
        "content_count": content_count or 0,
        "active_sessions": active_sessions or 0,
        "risk_alerts": risk_alerts or 0
    }


# ---------------------------------------------------------------------------
# RAG endpoints
# ---------------------------------------------------------------------------


@router.get("/search")
async def semantic_search(
    q: str = Query(..., min_length=2, description="Search query"),
    content_id: Optional[str] = None,
    top_k: int = Query(default=5, ge=1, le=20),
    current_user=Depends(get_current_user),
):
    """
    Semantic search across the pedagogical knowledge base.
    Returns the most relevant content chunks ranked by cosine similarity.
    """
    cid = UUID(content_id) if content_id else None
    results = await retrieve_relevant_chunks(q, content_id=cid, top_k=top_k)

    # Stringify UUIDs for JSON serialization
    for r in results:
        r["content_id"] = str(r["content_id"])

    return {"query": q, "results": results}


@router.post("/reembed")
async def reembed_all(current_user=Depends(get_current_user)):
    """
    Re-embeds all content items. Teacher/admin only.
    Useful after model upgrade or initial migration to RAG.
    """
    if current_user["role"] not in ("teacher", "admin"):
        raise HTTPException(status_code=403, detail="Teacher or admin access required")

    result = await reembed_all_content()
    return result


# ---------------------------------------------------------------------------
# Content Critic
# ---------------------------------------------------------------------------

@router.post("/{content_id}/review")
async def review_content_quality(content_id: UUID, current_user=Depends(get_current_user)):
    """
    Triggers the Content Critic Agent to review uploaded content.
    Returns constructive feedback on clarity, grade appropriateness, and completeness.
    Teacher only.
    """
    if current_user["role"] != "teacher":
        raise HTTPException(status_code=403, detail="Teacher access required")

    pool = await get_pool()
    content = await pool.fetchrow(
        "SELECT title, original_text, subject, grade_level FROM content_items WHERE id = $1 AND teacher_id = $2",
        content_id, current_user["id"],
    )
    if not content:
        raise HTTPException(status_code=404, detail="Content not found")

    from agents.content_critic.agent import review_content
    result = await review_content(
        content_id=content_id,
        title=content["title"],
        text=content["original_text"],
        subject=content["subject"] or "General",
        grade_level=content["grade_level"] or 3,
    )
    return result


@router.get("/{content_id}/reviews")
async def get_content_reviews(content_id: UUID, current_user=Depends(get_current_user)):
    """Get all reviews for a content item."""
    from agents.content_critic.agent import get_content_reviews as fetch_reviews
    reviews = await fetch_reviews(content_id)
    return reviews
